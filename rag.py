"""Tiny RAG engine: chunk the failure playbook and search it with BM25.

Pure Python on purpose: no API key, no model download, works offline, and the
ranking is easy to explain in a README.

Token counts are ESTIMATES (about 4 characters per token), not a real tokenizer.
"""
import html
import math
import os
import re
import sys
from collections import Counter

# ~225 words is roughly 300 tokens; 40 words of overlap keeps sentences that
# straddle a boundary retrievable.
CHUNK_WORDS = 225
OVERLAP_WORDS = 40
TOP_K = 3

_STOP = set("a an the is are was were of to in on for and or with at by from it "
            "this that as be".split())


def estimate_tokens(text: str) -> int:
    return math.ceil(len(text) / 4)


def _tokenize(text: str):
    return [w for w in re.findall(r"[a-z0-9]+", text.lower()) if w not in _STOP]

TEXT_EXTS = {".md", ".txt", ".csv", ".json", ".log"}


def _read_file(path):
    """Return the plain text of a file, or None if the type is not supported."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXTS:
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext in (".html", ".htm"):
        with open(path, encoding="utf-8", errors="replace") as f:
            raw = f.read()
        raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
        return html.unescape(re.sub(r"<[^>]+>", " ", raw))
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("PDF support needs:  python -m pip install pypdf")
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise RuntimeError("Word support needs:  python -m pip install python-docx")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return None

class PlaybookIndex:
    def __init__(self, folder, chunk_words=CHUNK_WORDS, overlap_words=OVERLAP_WORDS):
        self.chunks = []
        self.full_tokens = 0
        step = chunk_words - overlap_words
        for fn in sorted(os.listdir(folder)):
            path = os.path.join(folder, fn)
            if not os.path.isfile(path):
                continue
            try:
                text = _read_file(path)
            except Exception as e:  # one bad file must not break the server
                print(f"[rag] skipped {fn}: {e}", file=sys.stderr)
                continue
            if text is None:
                print(f"[rag] skipped {fn}: unsupported file type", file=sys.stderr)
                continue
            if not text.strip():
                print(f"[rag] skipped {fn}: no text found (scanned PDF? it needs OCR first)", file=sys.stderr)
                continue
            self.full_tokens += estimate_tokens(text)
            words, i, idx = text.split(), 0, 1
            while True:
                piece = " ".join(words[i:i + chunk_words])
                self.chunks.append({"source": fn, "part": idx, "text": piece,
                                    "tokens": estimate_tokens(piece)})
                if i + chunk_words >= len(words):
                    break
                i += step
                idx += 1
        # BM25 statistics (filename words are indexed too: they are good topic hints)
        self._doc_tokens = [_tokenize(c["source"].replace("_", " ") + " " + c["text"])
                            for c in self.chunks]
        self._avgdl = sum(map(len, self._doc_tokens)) / max(len(self._doc_tokens), 1)
        df = Counter()
        for toks in self._doc_tokens:
            df.update(set(toks))
        n = len(self._doc_tokens)
        self._idf = {t: math.log(1 + (n - d + 0.5) / (d + 0.5)) for t, d in df.items()}

    def search(self, query, k=TOP_K, k1=1.5, b=0.75):
        q = _tokenize(query)
        scored = []
        for chunk, toks in zip(self.chunks, self._doc_tokens):
            tf = Counter(toks)
            score = 0.0
            for term in q:
                if term in tf:
                    num = tf[term] * (k1 + 1)
                    den = tf[term] + k1 * (1 - b + b * len(toks) / self._avgdl)
                    score += self._idf.get(term, 0) * num / den
            if score > 0:
                scored.append((score, chunk))
        scored.sort(key=lambda s: s[0], reverse=True)
        return [c for _, c in scored[:k]]
